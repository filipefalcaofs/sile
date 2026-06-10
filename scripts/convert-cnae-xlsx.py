#!/usr/bin/env python3
"""Conversão one-shot do arquivo oficial CNAE-Subclasses 2.3 (IBGE/CONCLA) para CSV versionado.

Proveniência:
- Fonte oficial: IBGE/CONCLA — CNAE-Subclasses 2.3, arquivo
  "CNAE_Subclasses_2_3_Estrutura_Detalhada.xlsx" (cnae.ibge.gov.br / concla.ibge.gov.br),
  vigência 01/01/2019 (Resolução CONCLA nº 02/2018).
- Arquivo baixado em 2026-06-09 e versionado em docs/dados-oficiais/.
- A publicação oficial cita 1.332 subclasses; o arquivo de estrutura contém 1.331.
  A única subclasse ausente é a 9900-8/00 (Organismos internacionais e outras
  instituições extraterritoriais), confirmada existente na busca oficial da CONCLA,
  porém omitida na aba de estrutura (a classe 99.00-8 aparece sem linha de subclasse).
  Este script é fiel à fonte: NÃO inventa a linha ausente — a divergência é registrada
  no relatório do import (CnaeImportService, plano 02-04).

Saída: database/data/cnaes-subclasses-2-3.csv (1 cabeçalho + 1.331 linhas de dados),
com a hierarquia resolvida por carry-forward e códigos no formato oficial
(ex.: 0111-3/01, 01.11-3, 01.1, 01, A). A normalização para dígitos é
responsabilidade do CnaeImportService.
"""

import csv
import re
import sys
from pathlib import Path

REPO_ROOT = Path(__file__).resolve().parent.parent
SOURCE_XLSX = REPO_ROOT / 'docs/dados-oficiais/CNAE_Subclasses_2_3_Estrutura_Detalhada.xlsx'
TARGET_CSV = REPO_ROOT / 'database/data/cnaes-subclasses-2-3.csv'
SHEET_NAME = 'Estrutura Det. CNAE Subclass2.3'
SUBCLASS_PATTERN = re.compile(r'^\d{4}-\d/\d{2}$')
EXPECTED_SUBCLASSES = 1331

HEADER = [
    'section_code', 'section_description',
    'division_code', 'division_description',
    'group_code', 'group_description',
    'class_code', 'class_description',
    'subclass_code', 'subclass_description',
]


def clean(value) -> str:
    return '' if value is None else str(value).strip()


def main() -> int:
    from openpyxl import load_workbook

    workbook = load_workbook(SOURCE_XLSX, read_only=True, data_only=True)
    sheet = workbook[SHEET_NAME]

    section = division = group = cnae_class = ('', '')
    sections, divisions, groups, classes = set(), set(), set(), set()
    rows = []

    # Linhas 1-4 são título/cabeçalho; dados começam na linha 5.
    # Árvore com um único nível preenchido por linha: o contexto hierárquico
    # é resolvido por carry-forward (última seção/divisão/grupo/classe vista).
    for row in sheet.iter_rows(min_row=5, max_col=6, values_only=True):
        col_a, col_b, col_c, col_d, col_e, col_f = (clean(value) for value in row)

        if col_a:
            section = (col_a, col_f)
            sections.add(col_a)
        if col_b:
            division = (col_b, col_f)
            divisions.add(col_b)
        if col_c:
            group = (col_c, col_f)
            groups.add(col_c)
        if col_d:
            cnae_class = (col_d, col_f)
            classes.add(col_d)
        if col_e:
            rows.append([
                section[0], section[1],
                division[0], division[1],
                group[0], group[1],
                cnae_class[0], cnae_class[1],
                col_e, col_f,
            ])

    workbook.close()

    codes = [row[8] for row in rows]
    errors = []

    if len(rows) != EXPECTED_SUBCLASSES:
        errors.append(f'esperadas {EXPECTED_SUBCLASSES} subclasses, encontradas {len(rows)}')
    if len(set(codes)) != len(codes):
        duplicates = sorted({code for code in codes if codes.count(code) > 1})
        errors.append(f'subclass_code duplicado: {duplicates}')
    invalid_codes = [code for code in codes if not SUBCLASS_PATTERN.match(code)]
    if invalid_codes:
        errors.append(f'códigos fora do padrão DDDD-D/SS: {invalid_codes[:5]}')
    missing_descriptions = [row[8] for row in rows if row[9] == '']
    if missing_descriptions:
        errors.append(f'subclasses sem denominação: {missing_descriptions[:5]}')

    if errors:
        for error in errors:
            print(f'ERRO: {error}', file=sys.stderr)
        return 1

    TARGET_CSV.parent.mkdir(parents=True, exist_ok=True)
    with TARGET_CSV.open('w', newline='', encoding='utf-8') as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADER)
        writer.writerows(rows)

    print(f'Subclasses emitidas: {len(rows)}')
    print(f'Seções distintas: {len(sections)}')
    print(f'Divisões distintas: {len(divisions)}')
    print(f'Grupos distintos: {len(groups)}')
    print(f'Classes distintas: {len(classes)}')
    print(f'CSV gerado em: {TARGET_CSV.relative_to(REPO_ROOT)}')

    return 0


if __name__ == '__main__':
    sys.exit(main())
